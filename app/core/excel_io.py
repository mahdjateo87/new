from __future__ import annotations

import calendar
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.config import MONTHS_FR, exports_dir
from app.core.models import AppConfig, Assignment
from app.core.stats import (
    compute_month_stats,
    get_assignment,
    iter_service_columns,
    month_days,
    month_title,
    set_assignment,
    cell_key,
)


HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
SURCHARGE_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _style_header(cell) -> None:
    cell.font = Font(bold=True)
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def _style_cell(cell) -> None:
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def export_garde_excel(
    config: AppConfig,
    service_id: str,
    year: int,
    month: int,
    garde_data: dict[str, Any],
    output: Path | None = None,
) -> Path:
    service = config.service_by_id(service_id)
    if not service:
        raise ValueError(f"Service inconnu: {service_id}")

    wb = Workbook()
    ws = wb.active
    ws.title = month_title(month)[:31]

    ws.merge_cells("C3:D3")
    ws["C3"] = config.clinique
    ws["C3"].font = Font(bold=True, size=14)

    ws.merge_cells("B4:E5")
    ws["B4"] = f"Liste de garde des {service['nom'].lower()} - mois de {month_title(month)}"
    ws["B4"].font = Font(bold=True, size=12)

    ws.merge_cells("C6:D6")
    ws["C6"] = year
    ws["C6"].font = Font(bold=True)

    columns = iter_service_columns(service)
    start_row = 8
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + 2, end_column=1)
    ws.cell(start_row, 1, "DATE")
    _style_header(ws.cell(start_row, 1))

    col_index = 2
    for vacation in service.get("vacations", []):
        code = vacation["code"]
        cols = vacation.get("colonnes", [])
        if not cols:
            continue
        start_col = col_index
        label = vacation.get("label") or config.vacation_label(code)
        if len(cols) > 1:
            ws.merge_cells(
                start_row=start_row,
                start_column=start_col,
                end_row=start_row,
                end_column=start_col + len(cols) - 1,
            )
        ws.cell(start_row, start_col, label)
        _style_header(ws.cell(start_row, start_col))

        for sub_col in cols:
            if " - " in sub_col:
                groupe, poste = sub_col.split(" - ", 1)
                ws.cell(start_row + 1, col_index, groupe.strip())
                ws.cell(start_row + 2, col_index, poste.strip())
            else:
                ws.cell(start_row + 1, col_index, sub_col)
                ws.cell(start_row + 2, col_index, "")
            _style_header(ws.cell(start_row + 1, col_index))
            _style_header(ws.cell(start_row + 2, col_index))
            col_index += 1

    entries = garde_data.get("entries", {})
    row = start_row + 3
    for day_str in month_days(year, month):
        y, m, d = map(int, day_str.split("-"))
        ws.cell(row, 1, date(y, m, d))
        _style_cell(ws.cell(row, 1))
        ws.cell(row, 1).number_format = "DD/MM/YYYY"

        col_idx = 2
        for vacation_code, _column_name, key in columns:
            assignment = Assignment.from_dict(entries.get(day_str, {}).get(key))
            ws.cell(row, col_idx, assignment.display_name())
            _style_cell(ws.cell(row, col_idx))
            col_idx += 1
        row += 1

    stats_row = row + 2
    ws.cell(stats_row, 1, "COMPTAGE MENSUEL")
    ws.cell(stats_row, 1).font = Font(bold=True, size=12)

    stats = compute_month_stats(config, service_id, year, month)
    header_row = stats_row + 1
    headers = ["Personnel", "Total vacations", "Total heures", "Absences", "Retards", "Congés"]
    vacation_codes = sorted({code for row in stats for code in row["vacation_counts"]})
    headers.extend(config.vacation_label(code) for code in vacation_codes)
    headers.append("Surcharge")

    for idx, header in enumerate(headers, start=1):
        ws.cell(header_row, idx, header)
        _style_header(ws.cell(header_row, idx))

    data_row = header_row + 1
    for stat in stats:
        ws.cell(data_row, 1, stat["person"])
        ws.cell(data_row, 2, stat["total_vacations"])
        ws.cell(data_row, 3, round(stat["total_heures"], 1))
        ws.cell(data_row, 4, stat["absences"])
        ws.cell(data_row, 5, stat["retards"])
        ws.cell(data_row, 6, stat["conges"])
        for idx, code in enumerate(vacation_codes, start=7):
            ws.cell(data_row, idx, stat["vacation_counts"].get(code, 0))
        surcharge_cell = ws.cell(data_row, 7 + len(vacation_codes), "OUI" if stat["surcharge"] else "NON")
        if stat["surcharge"]:
            surcharge_cell.fill = SURCHARGE_FILL
        for col in range(1, len(headers) + 1):
            _style_cell(ws.cell(data_row, col))
        data_row += 1

    ws.cell(data_row + 1, 1, f"Chef de service : {config.chef_service}")
    ws.cell(data_row + 2, 1, f"Lieu : {config.lieu}")

    for col in range(1, col_index):
        ws.column_dimensions[get_column_letter(col)].width = 16

    if output is None:
        filename = f"garde_{service_id}_{year}_{month:02d}.xlsx"
        output = exports_dir() / filename
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def export_comptage_excel(
    config: AppConfig,
    service_id: str,
    year: int,
    month: int,
    output: Path | None = None,
) -> Path:
    service = config.service_by_id(service_id)
    if not service:
        raise ValueError(f"Service inconnu: {service_id}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Comptage"

    ws["A1"] = config.clinique
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Comptage des vacations - {service['nom']} - {month_title(month)} {year}"
    ws["A2"].font = Font(bold=True)

    stats = compute_month_stats(config, service_id, year, month)
    headers = ["Personnel", "Total vacations", "Total heures", "Absences", "Retards", "Congés"]
    vacation_codes = sorted({code for row in stats for code in row["vacation_counts"]})
    headers.extend(config.vacation_label(code) for code in vacation_codes)
    headers.append("Surcharge")

    row = 4
    for col, header in enumerate(headers, start=1):
        ws.cell(row, col, header)
        _style_header(ws.cell(row, col))

    row += 1
    for stat in stats:
        ws.cell(row, 1, stat["person"])
        ws.cell(row, 2, stat["total_vacations"])
        ws.cell(row, 3, round(stat["total_heures"], 1))
        ws.cell(row, 4, stat["absences"])
        ws.cell(row, 5, stat["retards"])
        ws.cell(row, 6, stat["conges"])
        for idx, code in enumerate(vacation_codes, start=7):
            ws.cell(row, idx, stat["vacation_counts"].get(code, 0))
        surcharge_cell = ws.cell(row, 7 + len(vacation_codes), "OUI" if stat["surcharge"] else "NON")
        if stat["surcharge"]:
            surcharge_cell.fill = SURCHARGE_FILL
        row += 1

    ws.cell(row + 1, 1, f"Seuil surcharge : {config.surcharge_heures} heures/mois")
    ws.cell(row + 2, 1, f"Chef de service : {config.chef_service}")

    if output is None:
        output = exports_dir() / f"comptage_{service_id}_{year}_{month:02d}.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def _normalize_name(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _guess_service_from_title(title: str, config: AppConfig) -> str | None:
    title_l = title.lower()
    for service in config.services:
        if service["nom"].lower() in title_l or service["id"] in title_l:
            return service["id"]
    if "instrument" in title_l:
        return "instrumentiste"
    if "hygi" in title_l:
        return "hygiene"
    if "stér" in title_l or "ster" in title_l:
        return "sterilisation"
    return None


def _extract_year(ws, title: str = "") -> int:
    for coord in ("C6", "B6", "E5", "C5", "D6"):
        val = ws[coord].value
        if isinstance(val, (int, float)) and 2000 <= int(val) <= 2100:
            return int(val)
        if isinstance(val, str) and val.strip().isdigit() and len(val.strip()) == 4:
            return int(val.strip())

    for row in range(1, 40):
        val = ws.cell(row, 1).value
        if isinstance(val, datetime):
            return val.year
        if isinstance(val, date):
            return val.year

    for token in title.split():
        if token.isdigit() and len(token) == 4:
            return int(token)
    return datetime.now().year


def import_garde_excel(path: Path, config: AppConfig, service_id: str | None = None) -> dict[str, Any]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active

    title = _normalize_name(ws["B4"].value or ws["A1"].value or ws["B1"].value)
    year = _extract_year(ws, title)

    month = None
    for idx, name in enumerate(MONTHS_FR, start=1):
        if name in (ws.title or "").lower() or name in title.lower():
            month = idx
            break
    if month is None:
        month = datetime.now().month

    if service_id is None:
        service_id = _guess_service_from_title(title, config) or config.services[0]["id"]

    service = config.service_by_id(service_id)
    if not service:
        raise ValueError("Service introuvable pour l'import.")

    columns = iter_service_columns(service)
    entries: dict[str, dict[str, Any]] = {}

    header_row = None
    for row in range(1, 20):
        val = ws.cell(row, 1).value
        if isinstance(val, str) and val.strip().upper() == "DATE":
            header_row = row
            break
    if header_row is None:
        raise ValueError("Impossible de trouver l'en-tête DATE dans le fichier Excel.")

    header_lines = 2
    if ws.cell(header_row + 1, 2).value and ws.cell(header_row + 2, 2).value:
        second = str(ws.cell(header_row + 1, 2).value).lower().strip()
        third = str(ws.cell(header_row + 2, 2).value).lower().strip()
        if second in ("mat", "chir") and ("salle" in third or third.replace(" ", "").isalnum()):
            header_lines = 3

    data_start = header_row + header_lines
    max_row = ws.max_row

    col_map: dict[int, str] = {}
    col_idx = 2
    for vacation_code, column_name, key in columns:
        col_map[col_idx] = key
        col_idx += 1

    for row in range(data_start, max_row + 1):
        day_val = ws.cell(row, 1).value
        if day_val is None:
            continue
        if isinstance(day_val, datetime):
            day_str = day_val.strftime("%Y-%m-%d")
        elif isinstance(day_val, date):
            day_str = day_val.strftime("%Y-%m-%d")
        else:
            continue

        if int(day_str.split("-")[1]) != month:
            continue

        for excel_col, key in col_map.items():
            raw = _normalize_name(ws.cell(row, excel_col).value)
            if not raw:
                continue
            assignment = Assignment(person=raw)
            if raw.startswith("[A]"):
                assignment.status = "absence"
                assignment.person = raw[3:].strip()
            elif raw.startswith("[R]"):
                assignment.status = "retard"
                assignment.person = raw[3:].strip()
            elif raw.startswith("[C]"):
                assignment.status = "conge"
                assignment.person = raw[3:].strip()
            elif "/" in raw:
                parts = [p.strip() for p in raw.split("/") if p.strip()]
                if len(parts) >= 2:
                    assignment.person = parts[0]
                    assignment.replacement = parts[-1]
                else:
                    assignment.person = raw
            else:
                assignment.person = raw
            set_assignment(entries, day_str, key, assignment)

    return {
        "service_id": service_id,
        "year": year,
        "month": month,
        "entries": entries,
        "notes": f"Importé depuis {path.name}",
    }
